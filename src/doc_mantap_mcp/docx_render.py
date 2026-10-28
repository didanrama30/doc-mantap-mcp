"""Render Markdown -> DOCX with python-docx."""

from __future__ import annotations

import io
import re
import urllib.request
from dataclasses import dataclass, replace
from pathlib import Path
from urllib.parse import urlparse

from bs4 import NavigableString, Tag
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Emu, Mm, Pt, RGBColor
from docx.text.run import Run

from . import highlight
from .charts import render_chart
from .common import CALLOUTS, ParsedDoc, labels, parse_markdown, prepare_output
from .pdf import PageSetup, _kpi_rows, page_setup
from .themes import COVER_STYLES, Theme, mix, on_color, shade, tint

MONO = "Consolas"
SYMBOL_FONT = "Segoe UI Symbol"
MAX_REMOTE_IMAGE = 15 * 1024 * 1024

_PPR = ["w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr",
        "w:widowControl", "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd", "w:tabs",
        "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
        "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
        "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
        "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
        "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
        "w:pPrChange"]
_RPR = ["w:rStyle", "w:rFonts", "w:b", "w:bCs", "w:i", "w:iCs", "w:caps", "w:smallCaps",
        "w:strike", "w:dstrike", "w:outline", "w:shadow", "w:emboss", "w:imprint",
        "w:noProof", "w:snapToGrid", "w:vanish", "w:webHidden", "w:color", "w:spacing", "w:w",
        "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight", "w:u", "w:effect", "w:bdr",
        "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang",
        "w:eastAsianLayout", "w:specVanish", "w:oMath"]
_TCPR = ["w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge", "w:tcBorders", "w:shd",
         "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark"]
_TBLPR = ["w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual", "w:tblStyleRowBandSize",
          "w:tblStyleColBandSize", "w:tblW", "w:jc", "w:tblCellSpacing", "w:tblInd",
          "w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
          "w:tblCaption", "w:tblDescription"]
_SETTINGS = ["w:defaultTabStop", "w:autoHyphenation", "w:consecutiveHyphenLimit",
             "w:hyphenationZone", "w:doNotHyphenateCaps", "w:showEnvelope", "w:summaryLength",
             "w:clickAndTypeStyle", "w:defaultTableStyle", "w:evenAndOddHeaders",
             "w:bookFoldRevPrinting", "w:bookFoldPrinting", "w:bookFoldPrintingSheets",
             "w:drawingGridHorizontalSpacing", "w:drawingGridVerticalSpacing",
             "w:displayHorizontalDrawingGridEvery", "w:displayVerticalDrawingGridEvery",
             "w:doNotUseMarginsForDrawingGridOrigin", "w:drawingGridHorizontalOrigin",
             "w:drawingGridVerticalOrigin", "w:doNotShadeFormData", "w:noPunctuationKerning",
             "w:characterSpacingControl", "w:printTwoOnOne", "w:strictFirstAndLastChars",
             "w:noLineBreaksAfter", "w:noLineBreaksBefore", "w:savePreviewPicture",
             "w:doNotValidateAgainstSchema", "w:saveInvalidXml", "w:ignoreMixedContent",
             "w:alwaysShowPlaceholderText", "w:doNotDemarcateInvalidXml", "w:saveXmlDataOnly",
             "w:useXSLTWhenSaving", "w:saveThroughXslt", "w:showXMLTags",
             "w:alwaysMergeEmptyNamespace", "w:updateFields", "w:hdrShapeDefaults",
             "w:footnotePr", "w:endnotePr", "w:compat", "w:docVars", "w:rsids", "m:mathPr",
             "w:attachedSchema", "w:themeFontLang", "w:clrSchemeMapping",
             "w:doNotIncludeSubdocsInStats", "w:doNotAutoCompressPictures", "w:forceUpgrade",
             "w:captions", "w:readModeInkLockDown", "w:smartTagType", "sl:schemaLibrary",
             "w:shapeDefaults", "w:doNotEmbedSmartTags", "w:decimalSymbol", "w:listSeparator"]


# --------------------------------------------------------------------------- #
# Low-level OOXML helpers
# --------------------------------------------------------------------------- #

def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.upper())


def _set(parent, tag: str, order: list[str], attrs: dict | None = None):
    """Replace/insert child `tag` keeping the schema order given by `order`."""
    old = parent.find(qn(tag))
    if old is not None:
        parent.remove(old)
    el = OxmlElement(tag)
    for k, v in (attrs or {}).items():
        el.set(qn(k), str(v))
    for later in order[order.index(tag) + 1:]:
        sibling = parent.find(qn(later))
        if sibling is not None:
            sibling.addprevious(el)
            return el
    parent.append(el)
    return el


def _font(target, name: str) -> None:
    """Set the Latin + complex-script font of a run or style."""
    target.font.name = name
    rpr = target.element.get_or_add_rPr() if hasattr(target, "element") else target._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = _set(rpr, "w:rFonts", _RPR)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), name)


def _letter_spacing(run, twentieths: int) -> None:
    _set(run._r.get_or_add_rPr(), "w:spacing", _RPR, {"w:val": twentieths})


def _run_shade(run, fill: str) -> None:
    _set(run._r.get_or_add_rPr(), "w:shd", _RPR,
         {"w:val": "clear", "w:color": "auto", "w:fill": fill})


def _para_border(paragraph, side: str, color: str, size: int, space: int = 4) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    bdr = ppr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = _set(ppr, "w:pBdr", _PPR)
    order = ["w:top", "w:left", "w:bottom", "w:right", "w:between", "w:bar"]
    _set(bdr, f"w:{side}", order, {"w:val": "single", "w:sz": size, "w:space": space,
                                    "w:color": color})


def _cell_fill(cell, fill: str) -> None:
    _set(cell._tc.get_or_add_tcPr(), "w:shd", _TCPR,
         {"w:val": "clear", "w:color": "auto", "w:fill": fill})


def _cell_margins(cell, top=None, bottom=None, left=None, right=None) -> None:
    """Margins in twips (1 cm = 567)."""
    tcpr = cell._tc.get_or_add_tcPr()
    mar = tcpr.find(qn("w:tcMar"))
    if mar is None:
        mar = _set(tcpr, "w:tcMar", _TCPR)
    order = ["w:top", "w:start", "w:left", "w:bottom", "w:end", "w:right"]
    for side, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        if value is not None:
            _set(mar, f"w:{side}", order, {"w:w": int(value), "w:type": "dxa"})


def _cell_borders(cell, **sides) -> None:
    """sides: top/left/bottom/right = (color, size) or None for no border."""
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = _set(tcpr, "w:tcBorders", _TCPR)
    order = ["w:top", "w:start", "w:left", "w:bottom", "w:end", "w:right", "w:insideH",
             "w:insideV", "w:tl2br", "w:tr2bl"]
    for side, spec in sides.items():
        attrs = {"w:val": "nil"} if spec is None else {
            "w:val": "single", "w:sz": spec[1], "w:space": 0, "w:color": spec[0]}
        _set(borders, f"w:{side}", order, attrs)


def _table_borders(table, **sides) -> None:
    tblpr = table._tbl.tblPr
    borders = _set(tblpr, "w:tblBorders", _TBLPR)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        spec = sides.get(side)
        el = OxmlElement(f"w:{side}")
        if spec is None:
            el.set(qn("w:val"), "nil")
        else:
            for k, v in {"w:val": "single", "w:sz": spec[1], "w:space": 0,
                         "w:color": spec[0]}.items():
                el.set(qn(k), str(v))
        borders.append(el)


def _table_cell_margins(table, top: int, bottom: int, left: int, right: int) -> None:
    mar = _set(table._tbl.tblPr, "w:tblCellMar", _TBLPR)
    for side, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        mar.append(el)


def _table_widths(table, widths_emu: list[int]) -> None:
    table.autofit = False
    total = sum(widths_emu)
    _set(table._tbl.tblPr, "w:tblW", _TBLPR, {"w:w": int(total / 635), "w:type": "dxa"})
    _set(table._tbl.tblPr, "w:tblInd", _TBLPR, {"w:w": 0, "w:type": "dxa"})
    for col, width in zip(table.columns, widths_emu):
        col.width = Emu(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_emu):
            cell.width = Emu(width)


def _row_flag(row, tag: str) -> None:
    trpr = row._tr.get_or_add_trPr()
    if trpr.find(qn(tag)) is None:
        trpr.append(OxmlElement(tag))


def _remove(paragraph) -> None:
    el = paragraph._p
    el.getparent().remove(el)


def _trim(paragraph) -> None:
    texts = paragraph._p.findall(".//" + qn("w:t"))
    if texts:
        texts[0].text = (texts[0].text or "").lstrip()
        texts[-1].text = (texts[-1].text or "").rstrip()


def _setting(document, tag: str, attrs: dict) -> None:
    _set(document.settings.element, tag, _SETTINGS, attrs)


@dataclass(frozen=True)
class Fmt:
    bold: bool = False
    italic: bool = False
    code: bool = False
    strike: bool = False
    underline: bool = False
    mark: bool = False
    sup: bool = False
    sub: bool = False
    color: str | None = None
    size: float | None = None
    font: str | None = None
    caps: bool = False


# --------------------------------------------------------------------------- #
# Builder
# --------------------------------------------------------------------------- #

class DocxBuilder:
    def __init__(self, theme: Theme, setup: PageSetup, lang: str):
        self.t = theme
        self.setup = setup
        self.lang = lang or "en"
        self.names = labels(self.lang)
        self.document = Document()
        self.width_emu = int(Cm(setup.text_width_cm))
        self._bookmark = 0
        self._anchors: dict[str, str] = {}
        self._fn_order: dict[str, int] = {}
        self._abstract: dict[bool, int] = {}
        self._pending_break = False
        self.doc: ParsedDoc | None = None

    # ---- setup -------------------------------------------------------------
    def setup_page(self) -> None:
        s = self.document.sections[0]
        s.page_width, s.page_height = Mm(self.setup.width_mm), Mm(self.setup.height_mm)
        if self.setup.width_mm > self.setup.height_mm:
            from docx.enum.section import WD_ORIENT
            s.orientation = WD_ORIENT.LANDSCAPE
        s.left_margin = s.right_margin = Mm(self.setup.margin_side_mm)
        s.top_margin, s.bottom_margin = Mm(self.setup.margin_top_mm), Mm(self.setup.margin_bottom_mm)
        s.header_distance = s.footer_distance = Mm(10)

    def setup_styles(self) -> None:
        t, styles = self.t, self.document.styles
        base = t.body_size_pt * (0.92 if self.setup.width_mm < 160 else 1)
        self.base = base
        normal = styles["Normal"]
        _font(normal, t.docx_body)
        normal.font.size = Pt(base)
        normal.font.color.rgb = _rgb(t.text)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.25
        # Proofing language (spell check + hyphenation)
        lang_code = {"id": "id-ID", "en": "en-US"}.get(self.lang.split("-")[0], self.lang)
        rpr_default = styles.element.find(qn("w:docDefaults") + "/" + qn("w:rPrDefault") + "/"
                                          + qn("w:rPr"))
        if rpr_default is not None:
            _set(rpr_default, "w:lang", _RPR, {"w:val": lang_code, "w:eastAsia": "en-US",
                                                 "w:bidi": "ar-SA"})

        sizes = {1: base * 2.05, 2: base * 1.45, 3: base * 1.18, 4: base * 0.85, 5: base,
                 6: base}
        for level, size in sizes.items():
            st = styles[f"Heading {level}"]
            _font(st, t.docx_heading)
            st.font.size = Pt(round(size * 2) / 2)
            st.font.bold = True
            st.font.italic = False
            st.font.color.rgb = _rgb(t.primary if level <= 3 else t.muted if level == 4
                                     else t.text)
            pf = st.paragraph_format
            pf.space_before = Pt({1: 20, 2: 18, 3: 12}.get(level, 10))
            pf.space_after = Pt({1: 10, 2: 6}.get(level, 4))
            pf.keep_with_next = True
            pf.line_spacing = 1.1
            if level == 4:
                st.font.all_caps = True
            ppr = st.element.get_or_add_pPr()
            for old in ppr.findall(qn("w:pBdr")):
                ppr.remove(old)
        h1, h2 = styles["Heading 1"], styles["Heading 2"]
        if t.heading_style == "bar":
            self._style_border(h1, "bottom", t.accent, 18, 6)
            self._style_border(h2, "left", t.accent, 24, 8)
        else:
            self._style_border(h1, "bottom", t.primary, 6, 6)

        for name, size, color, kw in (
            ("Caption", base * 0.8, t.muted, {}),
            ("Quote", base * 1.02, t.muted, {"italic": True}),
            ("Header", base * 0.7, t.muted, {}),
            ("Footer", base * 0.7, t.muted, {}),
            ("TOC Heading", base * 1.8, t.primary, {"bold": True}),
        ):
            st = styles[name]
            _font(st, t.docx_heading if name == "TOC Heading" else t.docx_body)
            st.font.size = Pt(round(size * 2) / 2)
            st.font.color.rgb = _rgb(color)
            st.font.italic = kw.get("italic", False)
            st.font.bold = kw.get("bold", False)
        styles["TOC Heading"].paragraph_format.space_after = Pt(14)
        styles["Caption"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        styles["Caption"].paragraph_format.space_before = Pt(4)
        quote = styles["Quote"]
        quote.paragraph_format.left_indent = Cm(0.5)
        quote.paragraph_format.space_before = Pt(6)
        quote.paragraph_format.space_after = Pt(6)
        self._style_border(quote, "left", t.accent, 18, 10)

        link = styles.add_style("Hyperlink", WD_STYLE_TYPE.CHARACTER)
        link.font.color.rgb = _rgb(shade(t.accent, 0.1))
        link.font.underline = True
        for level in (1, 2, 3):
            st = styles.add_style(f"toc {level}", WD_STYLE_TYPE.PARAGRAPH)
            st.base_style = styles["Normal"]
            st.font.bold = level == 1
            st.font.color.rgb = _rgb(t.primary if level == 1 else t.text if level == 2
                                     else t.muted)
            pf = st.paragraph_format
            pf.left_indent = Cm(0.6 * (level - 1))
            pf.space_before = Pt(8 if level == 1 else 2)
            pf.space_after = Pt(2)
            pf.tab_stops.add_tab_stop(Emu(self.width_emu), WD_TAB_ALIGNMENT.RIGHT,
                                      _tab_leader_dots())
        lp = styles["List Paragraph"].paragraph_format
        lp.space_after = Pt(3)

    @staticmethod
    def _style_border(style, side, color, size, space):
        ppr = style.element.get_or_add_pPr()
        bdr = ppr.find(qn("w:pBdr"))
        if bdr is None:
            bdr = _set(ppr, "w:pBdr", _PPR)
        _set(bdr, f"w:{side}", ["w:top", "w:left", "w:bottom", "w:right", "w:between",
                                "w:bar"],
             {"w:val": "single", "w:sz": size, "w:space": space, "w:color": color})

    # ---- runs & paragraphs -------------------------------------------------
    def para(self, container, style: str | None = None):
        p = container.add_paragraph(style=style) if style else container.add_paragraph()
        if self._pending_break:
            p.paragraph_format.page_break_before = True
            self._pending_break = False
        return p

    def run(self, paragraph, text: str, fmt: Fmt = Fmt(), parent=None) -> Run:
        r = OxmlElement("w:r")
        (parent if parent is not None else paragraph._p).append(r)
        run = Run(r, paragraph)
        run.text = text
        if fmt.bold:
            run.bold = True
        if fmt.italic:
            run.italic = True
        if fmt.strike:
            run.font.strike = True
        if fmt.underline:
            run.underline = True
        if fmt.sup:
            run.font.superscript = True
        if fmt.sub:
            run.font.subscript = True
        if fmt.caps:
            run.font.all_caps = True
        if fmt.code:
            _font(run, MONO)
            run.font.size = Pt(max(self.base - 1.5, 7))
            run.font.color.rgb = _rgb(shade(self.t.accent, 0.25))
            _run_shade(run, self.t.surface)
        if fmt.font:
            _font(run, fmt.font)
        if fmt.size:
            run.font.size = Pt(fmt.size)
        if fmt.color:
            run.font.color.rgb = _rgb(fmt.color)
        if fmt.mark:
            _run_shade(run, tint("FACC15", 0.55))
        return run

    def field(self, paragraph, instr: str, fmt: Fmt = Fmt(), placeholder: str = "") -> None:
        def char(kind):
            r = self.run(paragraph, "", fmt)
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            r._r.append(el)
        char("begin")
        r = self.run(paragraph, "", fmt)
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = f" {instr} "
        r._r.append(it)
        char("separate")
        self.run(paragraph, placeholder, fmt)
        char("end")

    def bookmark(self, paragraph, name: str) -> None:
        self._bookmark += 1
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(self._bookmark))
        start.set(qn("w:name"), name)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(self._bookmark))
        ppr = paragraph._p.pPr
        if ppr is not None:
            ppr.addnext(start)
        else:
            paragraph._p.insert(0, start)
        paragraph._p.append(end)

    def hyperlink(self, paragraph, url: str | None = None, anchor: str | None = None):
        h = OxmlElement("w:hyperlink")
        if url:
            h.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
        else:
            h.set(qn("w:anchor"), anchor)
        h.set(qn("w:history"), "1")
        paragraph._p.append(h)
        return h

    def inline(self, paragraph, node, fmt: Fmt = Fmt(), parent=None) -> None:
        """Add inline HTML content (text, emphasis, code, links, notes) as runs."""
        for child in node.children:
            if isinstance(child, NavigableString):
                text = re.sub(r"\s+", " ", str(child))
                if text:
                    self.run(paragraph, text, fmt, parent)
                continue
            if not isinstance(child, Tag):
                continue
            name = child.name
            if name == "br":
                self.run(paragraph, "", fmt, parent).add_break()
            elif name == "a" and "footnote-ref" in child.get("class", []):
                key = child.get("href", "").lstrip("#")
                num = self._fn_order.setdefault(key, len(self._fn_order) + 1)
                self.run(paragraph, str(num), replace(fmt, sup=True, bold=True,
                                                      color=self.t.accent), parent)
            elif name == "a":
                href = child.get("href", "")
                target = None
                if href.startswith("#"):
                    target = self.hyperlink(paragraph, anchor=self._anchors.get(href[1:], href[1:]))
                elif urlparse(href).scheme in ("http", "https", "mailto"):
                    target = self.hyperlink(paragraph, url=href)
                link_fmt = replace(fmt, color=shade(self.t.accent, 0.1), underline=True)
                self.inline(paragraph, child, link_fmt, target if target is not None else parent)
            elif name == "span" and "task-box" in child.get("class", []):
                done = "done" in child["class"]
                self.run(paragraph, "☑ " if done else "☐ ",
                         replace(fmt, font=SYMBOL_FONT,
                                 color=self.t.accent if done else self.t.muted), parent)
            elif name == "img":
                self.picture_run(paragraph, child, max_height_cm=1.2)
            elif name == "sup" and child.find("a", class_="footnote-ref"):
                self.inline(paragraph, child, fmt, parent)
            else:
                self.inline(paragraph, child, replace(
                    fmt,
                    bold=fmt.bold or name in ("strong", "b"),
                    italic=fmt.italic or name in ("em", "i", "cite"),
                    code=fmt.code or name in ("code", "kbd", "samp"),
                    strike=fmt.strike or name in ("del", "s", "strike"),
                    underline=fmt.underline or name in ("ins", "u"),
                    mark=fmt.mark or name == "mark",
                    sup=fmt.sup or name == "sup",
                    sub=fmt.sub or name == "sub",
                    color=self.t.primary if name in ("strong", "b") and fmt.color is None
                    and self.t.primary != self.t.text else fmt.color,
                ), parent)

    # ---- images -------------------------------------------------------------
    def _image_bytes(self, img: Tag) -> bytes | None:
        path = img.get("data-path")
        if path:
            p = Path(path)
            if p.suffix.lower() == ".svg" or not p.is_file():
                return None
            return p.read_bytes()
        src = img.get("src", "")
        if urlparse(src).scheme in ("http", "https"):
            try:
                req = urllib.request.Request(src, headers={"User-Agent": "doc-mantap-mcp"})
                with urllib.request.urlopen(req, timeout=15) as resp:
                    if not resp.headers.get_content_type().startswith("image/") or \
                            "svg" in resp.headers.get_content_type():
                        return None
                    return resp.read(MAX_REMOTE_IMAGE)
            except OSError:
                return None
        return None

    def picture_run(self, paragraph, img: Tag, max_height_cm: float | None = None) -> bool:
        data = self._image_bytes(img)
        if data is None:
            alt = img.get("alt") or img.get("src", "")
            self.run(paragraph, f"[{alt}]", Fmt(italic=True, color=self.t.muted))
            return False
        from PIL import Image
        with Image.open(io.BytesIO(data)) as im:
            w_px, h_px = im.size
            dpi = (im.info.get("dpi") or (96, 96))[0] or 96
        width = int(w_px / dpi * 914400)
        requested = img.get("width", "")
        if requested.endswith("%"):
            width = int(self.width_emu * float(requested[:-1]) / 100)
        elif requested.isdigit():
            width = int(int(requested) / 96 * 914400)
        width = min(width, self.width_emu)
        max_h = Cm(max_height_cm) if max_height_cm else Cm(self.setup.height_mm / 10 * 0.55)
        if h_px and width * h_px / w_px > max_h:
            width = int(max_h * w_px / h_px)
        self.run(paragraph, "").add_picture(io.BytesIO(data), width=Emu(width))
        return True

    # ---- blocks ---------------------------------------------------------------
    def blocks(self, container, nodes) -> None:
        for el in nodes:
            if isinstance(el, NavigableString):
                if el.strip():
                    p = self.para(container)
                    self.run(p, re.sub(r"\s+", " ", str(el)))
                    _trim(p)
                continue
            if not isinstance(el, Tag):
                continue
            self.block(container, el)

    def block(self, container, el: Tag) -> None:
        name, classes = el.name, el.get("class", [])
        if name in ("h1", "h2", "h3", "h4", "h5", "h6"):
            p = self.para(container, f"Heading {name[1]}")
            self.inline(p, el)
            _trim(p)
            if el.get("id"):
                self.bookmark(p, self._anchors.setdefault(el["id"], f"_Toc{len(self._anchors) + 1:05d}"))
        elif name == "p":
            p = self.para(container)
            if self.t.justify:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            self.inline(p, el)
            _trim(p)
        elif name in ("ul", "ol"):
            self.list(container, el, ordered=name == "ol")
        elif name == "table":
            self.table(container, el)
        elif name == "pre":
            self.code(container, el)
        elif name == "div" and "callout" in classes:
            self.callout(container, el)
        elif name == "div" and "dm-chart" in classes:
            self.chart(container, self.doc.charts[int(el["data-index"])])
        elif name == "div" and "dm-kpis" in classes:
            self.kpis(container, self.doc.kpis[int(el["data-index"])])
        elif name == "div" and "page-break" in classes:
            self._pending_break = True
        elif name == "figure":
            self.figure(container, el)
        elif name == "blockquote":
            for child in el.children:
                if isinstance(child, Tag) and child.name == "p":
                    q = self.para(container, "Quote")
                    self.inline(q, child, Fmt(italic=True))
                    _trim(q)
                elif isinstance(child, Tag):
                    self.block(container, child)
        elif name == "hr":
            p = self.para(container)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(12)
            _para_border(p, "bottom", self.t.border, 6, 1)
        elif name == "dl":
            for child in el.find_all(["dt", "dd"], recursive=False):
                p = self.para(container)
                if child.name == "dt":
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.keep_with_next = True
                    self.inline(p, child, Fmt(bold=True, color=self.t.primary))
                else:
                    p.paragraph_format.left_indent = Cm(0.6)
                    self.inline(p, child)
                _trim(p)
        elif name in ("div", "section", "article", "main", "details", "nav", "header",
                      "footer", "aside"):
            self.blocks(container, el.children)
        elif name in ("style", "script", "link", "meta"):
            return
        else:
            p = self.para(container)
            self.inline(p, el)
            _trim(p)

    # ---- lists ----------------------------------------------------------------
    def _abstract_id(self, ordered: bool) -> int:
        if ordered in self._abstract:
            return self._abstract[ordered]
        numbering = self.document.part.numbering_part.element
        ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
        aid = max(ids, default=0) + 1
        font, color = self.t.docx_body, self.t.accent
        levels = []
        for lvl in range(9):
            if ordered:
                fmt, text = ("decimal", "lowerLetter", "lowerRoman")[lvl % 3], f"%{lvl + 1}."
            else:
                fmt, text = "bullet", ("•", "–", "▪")[lvl % 3]
            bold = "<w:b/>" if ordered else ""
            levels.append(
                f'<w:lvl w:ilvl="{lvl}"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/>'
                f'<w:lvlText w:val="{text}"/><w:lvlJc w:val="left"/>'
                f'<w:pPr><w:ind w:left="{400 * (lvl + 1)}" w:hanging="{360 if ordered else 300}"/></w:pPr>'
                f'<w:rPr><w:rFonts w:ascii="{font}" w:hAnsi="{font}" w:cs="{font}"/>{bold}'
                f'<w:color w:val="{color}"/></w:rPr></w:lvl>')
        xml = (f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{aid}">'
               f'<w:multiLevelType w:val="hybridMultilevel"/>{"".join(levels)}</w:abstractNum>')
        abstract = parse_xml(xml)
        first_num = numbering.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(abstract)
        else:
            numbering.append(abstract)
        self._abstract[ordered] = aid
        return aid

    def _new_num(self, ordered: bool, level: int, start: int = 1) -> int:
        numbering = self.document.part.numbering_part.element
        num = numbering.add_num(self._abstract_id(ordered))
        num.add_lvlOverride(ilvl=level).add_startOverride(start)
        return num.numId

    def list(self, container, el: Tag, ordered: bool, level: int = 0) -> None:
        start = int(el.get("start", 1)) if str(el.get("start", "1")).isdigit() else 1
        num_id = self._new_num(ordered, level, start)
        for li in el.find_all("li", recursive=False):
            task = "task-list-item" in li.get("class", [])
            p = self.para(container, "List Paragraph")
            pf = p.paragraph_format
            if task:
                pf.left_indent = Cm(0.7 * (level + 1))
                pf.first_line_indent = Cm(-0.55)
            else:
                numpr = p._p.get_or_add_pPr().get_or_add_numPr()
                numpr.get_or_add_ilvl().val = level
                numpr.get_or_add_numId().val = num_id
            first = True
            for child in li.children:
                if isinstance(child, Tag) and child.name in ("ul", "ol"):
                    self.list(container, child, child.name == "ol", level + 1)
                elif isinstance(child, Tag) and child.name == "p":
                    if first:
                        self.inline(p, child)
                        first = False
                    else:
                        extra = self.para(container, "List Paragraph")
                        extra.paragraph_format.left_indent = Cm(0.7 * (level + 1))
                        self.inline(extra, child)
                        _trim(extra)
                elif isinstance(child, Tag) and child.name in ("pre", "table", "div",
                                                               "blockquote", "figure"):
                    self.block(container, child)
                elif isinstance(child, Tag):
                    self.inline(p, _wrap(child))
                    first = False
                elif str(child).strip():
                    self.inline(p, _wrap(child))
                    first = False
            _trim(p)

    # ---- tables -----------------------------------------------------------------
    def table(self, container, el: Tag) -> None:
        rows = [r.find_all(["th", "td"]) for r in el.find_all("tr")]
        rows = [r for r in rows if r]
        if not rows:
            return
        self._flush_break(container)
        ncols = max(len(r) for r in rows)
        lengths = [max(min(len(r[j].get_text(strip=True)), 40) if j < len(r) else 0
                       for r in rows) for j in range(ncols)]
        weights = [max(n, 4) ** 0.85 for n in lengths]
        widths = [int(self.width_emu * w / sum(weights)) for w in weights]
        table = container.add_table(rows=len(rows), cols=ncols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _table_widths(table, widths)
        t, lined = self.t, self.t.table_style == "lined"
        _table_borders(table, bottom=(t.text if lined else t.primary, 12),
                       insideH=(t.border, 4))
        _table_cell_margins(table, 70, 70, 115, 115)
        size = round(self.base * 0.9 * 2) / 2
        for i, cells in enumerate(rows):
            row = table.rows[i]
            header = all(c.name == "th" for c in cells)
            _row_flag(row, "w:cantSplit")
            if header and i == 0:
                _row_flag(row, "w:tblHeader")
            for j in range(ncols):
                cell = table.cell(i, j)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.1
                if j >= len(cells):
                    continue
                style = cells[j].get("style", "")
                if "right" in style:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                elif "center" in style:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if header:
                    color = t.text if lined else on_color(t.primary)
                    self.inline(p, cells[j], Fmt(bold=True, color=color, size=size))
                    if lined:
                        _cell_borders(cell, top=(t.text, 12), bottom=(t.text, 6))
                    else:
                        _cell_fill(cell, t.primary)
                else:
                    self.inline(p, cells[j], Fmt(size=size))
                    if not lined and i % 2 == 0:
                        _cell_fill(cell, t.zebra)
                _trim(p)
        self._spacer(container)

    def _spacer(self, container, pt: float = 6) -> None:
        p = container.add_paragraph()
        p.paragraph_format.space_after = Pt(pt)
        p.paragraph_format.line_spacing = Pt(2)

    def _flush_break(self, container) -> None:
        """Tables cannot carry 'page break before', so emit a tiny paragraph that does."""
        if self._pending_break:
            p = self.para(container)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(2)

    def _box(self, container, fill: str | None, left_border: tuple | None = None,
             margins=(160, 160, 220, 220)):
        """Single-cell table used for code blocks and callouts. Returns the cell."""
        self._flush_break(container)
        table = container.add_table(rows=1, cols=1)
        _table_widths(table, [self.width_emu])
        _table_borders(table)
        cell = table.cell(0, 0)
        if fill:
            _cell_fill(cell, fill)
        if left_border:
            _cell_borders(cell, left=left_border)
        _cell_margins(cell, *margins)
        _row_flag(table.rows[0], "w:cantSplit")
        return cell

    def code(self, container, pre: Tag) -> None:
        lang = pre.get("data-lang", "")
        cell = self._box(container, highlight.CODE_BG, margins=(170, 170, 260, 260))
        p = cell.paragraphs[0]
        label = highlight.code_label(lang)
        if label:
            self.run(p, label, Fmt(font=MONO, size=6.5, bold=True, color=highlight.CODE_LABEL))
            _letter_spacing(p.runs[-1], 30)
            p.paragraph_format.space_after = Pt(4)
            p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        code_fmt = Fmt(font=MONO, size=max(self.base - 2, 7.5), color=highlight.CODE_FG)
        for text, color, bold, italic in highlight.tokens(pre.get_text().rstrip("\n"), lang):
            parts = text.split("\n")
            for k, part in enumerate(parts):
                if k:
                    self.run(p, "", code_fmt).add_break()
                if part:
                    self.run(p, part, replace(code_fmt, color=color or highlight.CODE_FG,
                                              bold=bold, italic=italic))
        self._spacer(container)

    def callout(self, container, el: Tag) -> None:
        kind = el.get("data-kind", "note")
        color = CALLOUTS.get(kind, CALLOUTS["note"])[0]
        cell = self._box(container, tint(color, 0.92), left_border=(color, 24),
                         margins=(150, 150, 240, 240))
        title_tag = el.find("p", class_="callout-title")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        if title_tag is not None:
            self.run(p, title_tag.get_text(strip=True),
                     Fmt(bold=True, color=shade(color, 0.18), size=self.base * 0.95))
            title_tag.extract()
        outer = self.width_emu
        self.width_emu = int(outer - Cm(0.9))  # nested tables/images must fit the box
        try:
            self.blocks(cell, el.children)
        finally:
            self.width_emu = outer
        for para in cell.paragraphs[1:]:
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if cell.paragraphs and len(cell.paragraphs) > 1:
            cell.paragraphs[-1].paragraph_format.space_after = Pt(0)
        self._spacer(container, 8)

    def kpis(self, container, items: list[dict]) -> None:
        t = self.t
        self._flush_break(container)
        for row_items in _kpi_rows(items):
            n = len(row_items)
            table = container.add_table(rows=1, cols=n)
            _table_widths(table, [self.width_emu // n] * n)
            _table_borders(table)
            _set(table._tbl.tblPr, "w:tblCellSpacing", _TBLPR, {"w:w": 90, "w:type": "dxa"})
            for cell, k in zip(table.rows[0].cells, row_items):
                _cell_fill(cell, t.surface)
                _cell_margins(cell, 150, 170, 190, 150)
                bar = cell.paragraphs[0]
                bar_width = Cm(0.9)
                bar.paragraph_format.right_indent = Emu(max(self.width_emu // n - bar_width
                                                            - Cm(0.7), 0))
                bar.paragraph_format.space_after = Pt(5)
                bar.paragraph_format.line_spacing = Pt(3)
                _para_border(bar, "top", t.accent, 18, 0)
                self.run(bar, "", Fmt(size=2))
                p = cell.add_paragraph()
                p.paragraph_format.space_after = Pt(1)
                r = self.run(p, k["label"], Fmt(bold=True, size=6.8, color=t.muted, caps=True))
                _letter_spacing(r, 14)
                p = cell.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.0
                self.run(p, k["value"], Fmt(bold=True, size=round(self.base * 1.75),
                                            color=t.primary, font=t.docx_heading))
                if k["delta"]:
                    tone = {"good": ("047857", "D1FAE5"), "bad": ("B91C1C", "FEE2E2")}.get(
                        k["tone"], (t.muted, "FFFFFF"))
                    arrow = {"up": "▲ ", "down": "▼ "}.get(k["trend"], "")
                    if k["delta"][:1] in "▲▼↑↓":
                        arrow = ""
                    p = cell.add_paragraph()
                    p.paragraph_format.space_after = Pt(1)
                    r = self.run(p, f" {arrow}{k['delta']} ", Fmt(bold=True, size=7.5,
                                                                  color=tone[0]))
                    _run_shade(r, tone[1])
                if k["note"]:
                    p = cell.add_paragraph()
                    p.paragraph_format.space_after = Pt(0)
                    self.run(p, k["note"], Fmt(size=7, color=t.muted))
            self._spacer(container, 4)
        self._spacer(container, 4)

    def chart(self, container, spec: dict) -> None:
        t = self.t
        if spec["title"]:
            p = self.para(container)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(1)
            self.run(p, str(spec["title"]), Fmt(bold=True, color=t.text, font=t.docx_heading,
                                                size=self.base))
        if spec["subtitle"]:
            p = self.para(container)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(2)
            self.run(p, str(spec["subtitle"]), Fmt(color=t.muted, size=self.base * 0.82))
        png = render_chart(spec, t, "png", self.setup.text_width_cm, self.lang)
        from PIL import Image
        with Image.open(io.BytesIO(png)) as im:
            width = min(int(im.size[0] / 220 * 914400), self.width_emu)
        p = self.para(container)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = bool(spec["caption"])
        self.run(p, "").add_picture(io.BytesIO(png), width=Emu(width))
        if spec["caption"]:
            cap = self.para(container, "Caption")
            cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
            self.run(cap, str(spec["caption"]))
        else:
            p.paragraph_format.space_after = Pt(10)

    def figure(self, container, el: Tag) -> None:
        img = el.find("img")
        if img is None:
            self.blocks(container, el.children)
            return
        p = self.para(container)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = el.find("figcaption")
        p.paragraph_format.keep_with_next = cap is not None
        self.picture_run(p, img)
        if cap is not None:
            c = self.para(container, "Caption")
            self.inline(c, cap)
            _trim(c)

    # ---- document furniture --------------------------------------------------
    @staticmethod
    def _next_para(cell):
        """The cell's first paragraph while it is still empty, otherwise a new one."""
        paragraphs = cell.paragraphs
        if len(paragraphs) == 1 and not paragraphs[0]._p.findall(qn("w:r")):
            return paragraphs[0]
        return cell.add_paragraph()

    def cover(self, style: str, title: str, subtitle, author, date, organization,
              document_type, logo: Path | None) -> None:
        t, s, doc, W = self.t, self.setup, self.document, self.width_emu
        V = WD_CELL_VERTICAL_ALIGNMENT
        text_h = Cm((s.height_mm - s.margin_top_mm - s.margin_bottom_mm) / 10 - 1.2)
        light, kicker_dark = on_color(t.primary), tint(t.accent, 0.45)
        title_pt = 30 if len(title) <= 30 else 25 if len(title) <= 60 else 21
        meta = [(k, v) for k, v in ((self.names["prepared_by"], author),
                                    (self.names["date"], date),
                                    (self.names["organization"], organization)) if v]

        def para(cell, center=False):
            p = self._next_para(cell)
            p.paragraph_format.space_after = Pt(0)
            if center:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            return p

        def rule_bar(cell, color, width_cm, center=False):
            p = para(cell, center)
            p.paragraph_format.line_spacing = Pt(2)
            inner = (cell.width or W) - Cm(2.4)
            if center:
                side = Emu(max((inner - Cm(width_cm)) // 2, 0))
                p.paragraph_format.left_indent = p.paragraph_format.right_indent = side
            else:
                p.paragraph_format.right_indent = Emu(max(inner - Cm(width_cm), 0))
            _para_border(p, "top", color, 24, 0)
            self.run(p, "", Fmt(size=2))
            return p

        def head(cell, dark=False, center=False, rule=None, rule_first=False):
            if rule and rule_first:
                rule_bar(cell, rule, 2.6).paragraph_format.space_after = Pt(18)
            if document_type:
                p = para(cell, center)
                p.paragraph_format.space_after = Pt(6)
                color = kicker_dark if dark else (t.muted if center else t.accent)
                _letter_spacing(self.run(p, document_type, Fmt(bold=True, size=8.5, caps=True,
                                                                color=color)), 40)
            p = para(cell, center)
            p.paragraph_format.line_spacing = 1.0
            self.run(p, title, Fmt(bold=True, size=title_pt - (4 if center else 0),
                                   font=t.docx_heading, color=light if dark else t.primary))
            if rule and not rule_first:
                rule_bar(cell, rule, 1.6 if center else 2.4, center) \
                    .paragraph_format.space_before = Pt(14)
            if subtitle:
                p = para(cell, center)
                p.paragraph_format.space_before = Pt(10)
                self.run(p, subtitle, Fmt(size=13, italic=center and not dark,
                                          color=mix(light, t.primary, 0.2) if dark else t.muted))

        def meta_block(cell, dark=False, center=False):
            if not meta:
                return
            if center:
                self.run(para(cell, True), "  ·  ".join(v for _, v in meta),
                         Fmt(size=9.5, color=t.muted, bold=True))
                return
            for kind in ("label", "value"):
                p = para(cell)
                p.paragraph_format.space_after = Pt(2)
                for i in range(1, len(meta)):
                    p.paragraph_format.tab_stops.add_tab_stop(Cm(5.0 * i))
                for i, (label, value) in enumerate(meta):
                    if i:
                        self.run(p, "\t")
                    if kind == "label":
                        _letter_spacing(self.run(p, label, Fmt(
                            bold=True, size=6.8, caps=True,
                            color=mix(light, t.primary, 0.35) if dark else t.muted)), 20)
                    else:
                        self.run(p, value, Fmt(bold=True, size=10.5,
                                               color=light if dark else t.text))

        def logo_into(cell, center=False, space_after=40):
            if logo is not None:
                p = para(cell, center)
                p.paragraph_format.space_after = Pt(space_after)
                self.picture_run(p, Tag(name="img", attrs={"data-path": str(logo)}),
                                 max_height_cm=1.5)

        def card(rows, fill):
            table = doc.add_table(rows=len(rows), cols=1)
            _table_widths(table, [W])
            _table_borders(table)
            for row, (height, valign) in zip(table.rows, rows):
                row.height = Emu(int(height))
                row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
                cell = row.cells[0]
                cell.vertical_alignment = valign
                if fill:
                    _cell_fill(cell, fill)
                _cell_margins(cell, 680, 680, 680, 680)
            return table

        if style == "band":
            table = card([(Cm(12.5), V.BOTTOM)], t.primary)
            head(table.cell(0, 0), dark=True)
            bar = doc.add_paragraph()
            bar.paragraph_format.left_indent = Cm(1.2)
            bar.paragraph_format.right_indent = Emu(max(W - Cm(4.4), 0))
            bar.paragraph_format.line_spacing = Pt(2)
            _para_border(bar, "top", t.accent, 48, 0)
            self.run(bar, "", Fmt(size=2))
            rest = card([(Cm(9.0), V.TOP)], None).cell(0, 0)
            _cell_margins(rest, 500, 0, 680, 680)
            meta_block(rest)
            if logo is not None:
                para(rest).paragraph_format.space_after = Pt(60)
                logo_into(rest)
        elif style == "full":
            top_h, mid_h = Cm(5.5), Cm(12.0)
            table = card([(top_h, V.TOP), (mid_h, V.TOP), (text_h - top_h - mid_h, V.BOTTOM)],
                         t.primary)
            logo_into(table.cell(0, 0), space_after=0)
            head(table.cell(1, 0), dark=True, rule=kicker_dark, rule_first=True)
            meta_block(table.cell(2, 0), dark=True)
        elif style == "classic":
            table = doc.add_table(rows=2, cols=3)
            spine, gap = int(Cm(0.55)), int(Cm(0.3))
            _table_widths(table, [spine, gap, W - spine - gap])
            _table_borders(table)
            top_h = Cm(15.0)
            for row, height in zip(table.rows, (top_h, text_h - top_h)):
                row.height = Emu(int(height))
                row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
            _cell_fill(table.cell(0, 0).merge(table.cell(1, 0)), t.primary)
            table.cell(0, 1).merge(table.cell(1, 1))
            content, foot = table.cell(0, 2), table.cell(1, 2)
            for c in (content, foot):
                _cell_borders(c, left=(t.accent, 8))
                _cell_margins(c, 1100, 900, 760, 200)
            foot.vertical_alignment = V.BOTTOM
            logo_into(content, space_after=70)
            head(content, rule=t.accent)
            meta_block(foot)
        else:  # minimal
            table = card([(Cm(15.0), V.CENTER), (text_h - Cm(15.0), V.BOTTOM)], None)
            top = table.cell(0, 0)
            logo_into(top, center=True)
            head(top, center=True, rule=t.accent)
            meta_block(table.cell(1, 0), center=True)
        self._pending_break = True

    def header_footer(self, title: str, footer_left: str, has_cover: bool,
                      section_style: str | None) -> None:
        t = self.t
        sec = self.document.sections[0]
        sec.different_first_page_header_footer = True
        right = Emu(self.width_emu)

        def prepare(p):
            for pos in (Emu(int(Cm(8.25))), Emu(int(Cm(16.5)))):
                p.paragraph_format.tab_stops.add_tab_stop(pos, WD_TAB_ALIGNMENT.CLEAR)
            p.paragraph_format.tab_stops.add_tab_stop(right, WD_TAB_ALIGNMENT.RIGHT)
            p.paragraph_format.space_after = Pt(0)

        hp = sec.header.paragraphs[0]
        prepare(hp)
        small = Fmt(size=7, color=t.muted)
        if title:
            r = self.run(hp, title, replace(small, caps=True))
            _letter_spacing(r, 16)
        if section_style:
            self.run(hp, "\t", small)
            self.field(hp, f'STYLEREF "{section_style}"', small)

        def footer(p):
            prepare(p)
            if footer_left:
                self.run(p, footer_left, small)
            self.run(p, "\t", small)
            num = Fmt(size=7.5, bold=True, color=t.primary)
            self.field(p, "PAGE", num, "1")
            self.run(p, " / ", num)
            self.field(p, "NUMPAGES", num, "1")

        footer(sec.footer.paragraphs[0])
        if not has_cover:
            footer(sec.first_page_footer.paragraphs[0])

    def toc(self, doc: ParsedDoc) -> None:
        heads = [h for h in doc.headings if h.id]
        if not heads:
            return
        top = min(h.level for h in heads)
        depth = 3 if sum(1 for h in heads if h.level <= top + 1) < 12 else 2
        title = self.para(self.document, "TOC Heading")
        self.run(title, self.names["contents"])
        entries = [h for h in heads if h.level - top < depth]
        last = len(entries) - 1
        for i, h in enumerate(entries):
            name = self._anchors.setdefault(h.id, f"_Toc{len(self._anchors) + 1:05d}")
            p = self.document.add_paragraph(style=f"toc {h.level - top + 1}")
            if i == 0:
                self.field_begin(p, f'TOC \\o "{top}-{top + depth - 1}" \\h \\z \\u')
            link = self.hyperlink(p, anchor=name)
            self.run(p, h.text, Fmt(), parent=link)
            if i == last:
                self.field_end(p)
        self._pending_break = True

    def field_begin(self, paragraph, instr: str) -> None:
        r = self.run(paragraph, "")
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), "begin")
        r._r.append(el)
        r = self.run(paragraph, "")
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = f" {instr} "
        r._r.append(it)
        r = self.run(paragraph, "")
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), "separate")
        r._r.append(el)

    def field_end(self, paragraph) -> None:
        r = self.run(paragraph, "")
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), "end")
        r._r.append(el)

    def notes(self, doc: ParsedDoc) -> None:
        if not self._fn_order:
            return
        p = self.para(self.document)
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.keep_with_next = True
        _para_border(p, "top", self.t.border, 6, 8)
        self.run(p, self.names["notes"], Fmt(bold=True, size=self.base * 0.9,
                                             color=self.t.primary))
        for key, num in sorted(self._fn_order.items(), key=lambda kv: kv[1]):
            note = doc.footnotes.get(key)
            q = self.para(self.document)
            q.paragraph_format.space_after = Pt(2)
            q.paragraph_format.left_indent = Cm(0.5)
            q.paragraph_format.first_line_indent = Cm(-0.5)
            self.run(q, f"{num}.\t", Fmt(bold=True, size=self.base * 0.8, color=self.t.accent))
            q.paragraph_format.tab_stops.add_tab_stop(Cm(0.5))
            if note is not None:
                for k, block in enumerate(note.find_all("p", recursive=False) or [note]):
                    if k:
                        self.run(q, " ")
                    self.inline(q, block, Fmt(size=self.base * 0.8, color=self.t.muted))
            _trim(q)


def _tab_leader_dots():
    from docx.enum.text import WD_TAB_LEADER
    return WD_TAB_LEADER.DOTS


def _wrap(node) -> Tag:
    """Wrap a single node in a span so it can go through `inline`."""
    span = Tag(name="span")
    span.append(node.__copy__() if isinstance(node, Tag) else NavigableString(str(node)))
    return span


def render_docx(content: str, output_path: str, theme: Theme, title: str | None = None,
                subtitle: str | None = None, author: str | None = None,
                date: str | None = None, *, organization: str | None = None,
                document_type: str | None = None, logo: str | None = None,
                cover_style: str | None = None, toc: bool = False, page_size: str = "A4",
                orientation: str = "portrait", lang: str = "en",
                base_dir: str | None = None) -> Path:
    out = prepare_output(output_path, ".docx")
    base = Path(base_dir).expanduser().resolve() if base_dir else out.parent
    style = (cover_style or theme.cover).lower()
    if style not in COVER_STYLES:
        raise ValueError(f"Unknown cover_style '{cover_style}'. Use: {', '.join(COVER_STYLES)}.")
    logo_path = None
    if logo:
        logo_path = Path(logo).expanduser()
        logo_path = (logo_path if logo_path.is_absolute() else base / logo_path).resolve()
        if not logo_path.is_file():
            raise ValueError(f"Logo file not found: {logo_path}")

    doc = parse_markdown(content, lang=lang, base_dir=base)
    b = DocxBuilder(theme, page_setup(page_size, orientation), lang)
    b.doc = doc
    b.setup_page()
    b.setup_styles()
    d = b.document
    first_h1 = next((h.text for h in doc.headings if h.level == 1), "")
    section_style = ("Heading 1" if any(h.level == 1 for h in doc.headings) else
                     "Heading 2" if any(h.level == 2 for h in doc.headings) else None)
    b.header_footer(title or first_h1, organization or author or "", bool(title), section_style)
    props = d.core_properties
    props.title = title or first_h1 or ""
    props.subject = subtitle or ""
    props.author = author or ""
    props.comments = "Created with doc-mantap-mcp"
    props.language = lang
    if title:
        b.cover(style, title, subtitle, author, date, organization, document_type, logo_path)
    if toc:
        b.toc(doc)
        _setting(d, "w:updateFields", {"w:val": "true"})
    if theme.justify:
        _setting(d, "w:autoHyphenation", {"w:val": "true"})
    zoom = d.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")  # required by the schema, missing in the template
    b.blocks(d, list(doc.soup.children))
    b.notes(doc)
    d.save(str(out))
    return out
