"""doc-mantap-mcp server."""

import functools
import json
import os
import zipfile
from typing import Literal

from docx.oxml.ns import qn
from lxml import etree
from mcp.server.mcpserver import Image, MCPServer
from mcp.server.mcpserver.exceptions import ToolError
from mcp.server.mcpserver.resources.templates import ResourceSecurity
from mcp.server.streamable_http import TransportSecuritySettings
from mcp_types import ToolAnnotations

from .docx_render import render_docx
from .docx_inspect import outline_docx, read_docx
from .docx_patch import (fill_template, find_text, insert_image, patch_docx,
                         replace_text_all)
from .doctor import environment_report
from .guide import guide
from .pdf import render_pdf
from .pdf_read import read_pdf_text
from .preview import render_pages
from .themes import THEMES, get_theme
from .xlsx_render import render_xlsx
from .xlsx_spec import Sheet

mcp = MCPServer(
    "doc-mantap-mcp",
    instructions=(
        "Turns AI-written content into beautiful files.\n"
        "- Documents: write Markdown, then call markdown_to_pdf (most polished) or "
        "markdown_to_docx (editable Word). Pass `title` (plus subtitle, document_type, author, "
        "date, organization, logo) for a designed cover; don't repeat the title as a heading. "
        "Use toc=true for long documents and lang='id' for Indonesian text.\n"
        "- Rich blocks: callouts (`> [!TIP]`), ```kpi cards, ```chart blocks "
        "(bar/line/area/pie/donut/scatter), footnotes, task lists, images, page breaks "
        "(\\pagebreak).\n"
        "- Spreadsheets: create_excel builds designed workbooks with real Excel tables, "
        "formulas, KPI cards, conditional formatting and native charts.\n"
        "- Existing files: docx_read/docx_outline inspect Word templates, docx_patch makes "
        "anchored edits, docx_find/docx_replace_all search templates, and "
        "docx_fill_template fills placeholders in single or batch mode. pdf_read_text extracts "
        "PDF text. markdown_to_docx supports mode='append' with target_path. Use doctor to "
        "check dependencies; docx://{path}#outline exposes document structure as a resource.\n"
        "Call design_guide once for the full syntax and design rules, list_themes for themes, "
        "and preview_document to look at the result before handing it over."
    ),
)

_WRITE = {"read_only_hint": False, "destructive_hint": True, "idempotent_hint": True}


def _explain(func):
    """Send anticipated errors (bad spec, missing library, file problems) to the model.

    The SDK hides the text of any exception that is not a ToolError, which would leave the
    model unable to fix its input.
    """
    @functools.wraps(func)
    def wrapper(*args, **kwargs):
        try:
            return func(*args, **kwargs)
        except ToolError:
            raise
        except (ValueError, RuntimeError, OSError, KeyError, TypeError, IndexError,
                zipfile.BadZipFile, etree.XMLSyntaxError) as exc:
            raise ToolError(str(exc)) from exc
    return wrapper


def _docx_counts(path: str) -> tuple[int, int]:
    """Count body paragraphs/tables directly so DOCM is supported too."""
    with zipfile.ZipFile(path) as package:
        root = etree.fromstring(package.read("word/document.xml"))
    return len(list(root.iter(qn("w:p")))), len(list(root.iter(qn("w:tbl"))))


@mcp.tool(annotations=ToolAnnotations(title="List themes", read_only_hint=True))
def list_themes() -> list[dict]:
    """List the available themes with their colors, fonts and default cover style."""
    return [t.summary() for t in THEMES.values()]


@mcp.tool(annotations=ToolAnnotations(title="Design guide", read_only_hint=True))
def design_guide(topic: Literal["documents", "excel", "all"] = "all") -> str:
    """Syntax reference (callouts, KPI cards, charts, Excel spec) and design rules.

    Read it before writing a long document or a workbook.
    """
    return guide(topic)


@mcp.tool(annotations=ToolAnnotations(title="Read Word document", read_only_hint=True))
@_explain
def docx_read(path: str, include_runs: bool = True,
              include_headers_footers: bool = True) -> dict:
    """Read DOCX text, paragraph styles, run formatting, tables and document furniture as JSON.

    This reads OOXML directly through python-docx; LibreOffice is not required.
    """
    return read_docx(path, include_runs, include_headers_footers)


@mcp.tool(annotations=ToolAnnotations(title="Word document outline", read_only_hint=True))
@_explain
def docx_outline(path: str) -> dict:
    """Return compact heading hierarchy, table dimensions and detected placeholders."""
    return outline_docx(path)


@mcp.tool(annotations=ToolAnnotations(title="Patch Word document", read_only_hint=False,
                                     destructive_hint=True, idempotent_hint=True))
@_explain
def docx_patch(path: str | None = None, operations: list[dict] | None = None,
               output_path: str | None = None, backup: bool = True,
               overwrite: bool = False, idempotent: bool = True,
               batch: list[dict] | None = None) -> dict:
    """Patch existing DOCX/DOCM in OOXML without regenerating its structure.

    Operations: replace_text(anchor, new_text), insert_after(anchor, paragraphs),
    fill_table_cell(table, row, col, content), append_to_cell(...), or
    insert_code_block(anchor, text). Default output is a new .patched file. To edit
    in place, set output_path to path; an unused .bak is required automatically.
    """
    if batch is not None:
        results = []
        for item in batch:
            results.append(patch_docx(item["path"], item["operations"],
                                      item.get("output_path"), backup, overwrite, idempotent))
        return {"count": len(results), "results": results}
    if not path or operations is None:
        raise ValueError("path and operations are required unless batch is provided")
    return patch_docx(path, operations, output_path, backup, overwrite, idempotent)


@mcp.tool(annotations=ToolAnnotations(title="Find Word text", read_only_hint=True))
@_explain
def docx_find(path: str, query: str, regex: bool = False,
              case_sensitive: bool = True) -> dict:
    """Find text across body paragraphs, tables, headers and footers."""
    return find_text(path, query, regex, case_sensitive)


@mcp.tool(annotations=ToolAnnotations(title="Replace Word text", read_only_hint=False,
                                     destructive_hint=True, idempotent_hint=True))
@_explain
def docx_replace_all(path: str, find: str, replace: str, regex: bool = False,
                     case_sensitive: bool = True, output_path: str | None = None,
                     backup: bool = True, overwrite: bool = False) -> dict:
    """Run-aware replacement across body, tables, headers and footers."""
    return replace_text_all(path, find, replace, regex, case_sensitive,
                            output_path, backup, overwrite)


@mcp.tool(annotations=ToolAnnotations(title="Fill Word template", read_only_hint=False,
                                     destructive_hint=True, idempotent_hint=True))
@_explain
def docx_fill_template(path: str | None = None, data: dict | None = None,
                       output_path: str | None = None, backup: bool = True,
                       overwrite: bool = False, batch: list[dict] | None = None) -> dict:
    """Fill {{key}} placeholders, including table/header/footer text; supports batches.

    Batch items: {"path": "...docx", "data": {"key": "value"}, "output_path": "..."}.
    Repeating a completed fill is a no-op rather than duplicating text.
    """
    if batch is not None:
        results = []
        for item in batch:
            item_data = item.get("data", data)
            if item_data is None:
                raise ValueError("Each batch item needs data, or pass shared data")
            results.append(fill_template(item["path"], item_data, item.get("output_path"),
                                         backup, overwrite))
        return {"count": len(results), "results": results}
    if not path or data is None:
        raise ValueError("path and data are required unless batch is provided")
    return fill_template(path, data, output_path, backup, overwrite)


@mcp.tool(annotations=ToolAnnotations(title="Environment doctor", read_only_hint=True))
def doctor() -> dict:
    """Check Python/native dependencies, Office tools, OCR languages and fonts."""
    return environment_report()


@mcp.resource("docx://{+path}#outline", name="DOCX outline", mime_type="application/json",
              description="Heading and table outline for a local DOCX/DOCM file.",
              security=ResourceSecurity(reject_absolute_paths=False))
@_explain
def docx_outline_resource(path: str) -> str:
    """Expose the same compact outline through a docx:// resource URI."""
    return json.dumps(outline_docx(path), ensure_ascii=False)


@mcp.tool(annotations=ToolAnnotations(title="Insert image in Word", read_only_hint=False,
                                     destructive_hint=True, idempotent_hint=False))
@_explain
def docx_insert_image(path: str, image_path: str, anchor: str | None = None,
                      position: Literal["after", "before", "replace", "append_cell"] = "after",
                      width_cm: float | None = None, caption: str | None = None,
                      alignment: Literal["left", "center", "right"] = "center",
                      table: int = 0, row: int | None = None, col: int | None = None,
                      output_path: str | None = None, backup: bool = True,
                      overwrite: bool = False) -> dict:
    """Insert image before/after a paragraph, replace it, or append to a table cell.

    Preserves existing package parts. Defaults to a new .patched file; explicit in-place
    edits create a .bak first. Table row/col are zero-based logical cell coordinates.
    """
    return insert_image(path, image_path, anchor, position, width_cm, caption, alignment,
                        table, row, col, output_path, backup, overwrite)


@mcp.tool(annotations=ToolAnnotations(title="Read PDF text", read_only_hint=True))
@_explain
def pdf_read_text(path: str, ocr: bool = True, language: str = "eng+ind") -> dict:
    """Extract native PDF text per page; OCR blank/scanned pages with Tesseract if installed."""
    return read_pdf_text(path, ocr, language)


@mcp.tool(annotations=ToolAnnotations(title="Markdown to PDF", **_WRITE))
@_explain
def markdown_to_pdf(content: str, output_path: str, theme: str = "modern",
                    title: str | None = None, subtitle: str | None = None,
                    author: str | None = None, date: str | None = None,
                    organization: str | None = None, document_type: str | None = None,
                    logo: str | None = None, cover_style: str | None = None,
                    toc: bool = False, page_size: str = "A4", orientation: str = "portrait",
                    lang: str = "en", primary_color: str | None = None,
                    accent_color: str | None = None, base_dir: str | None = None) -> str:
    """Convert Markdown into a designed PDF (cover, running header, page numbers).

    Supports callouts (> [!TIP]), ```kpi and ```chart blocks, footnotes, task lists and
    syntax-highlighted code. See design_guide for the syntax.

    Args:
        content: Document body in Markdown.
        output_path: Output file path, e.g. ~/Documents/report.pdf.
        theme: modern, corporate, academic, fresh, elegant, minimal, startup or warm.
        title: Adds a cover page with this title.
        subtitle: Cover subtitle.
        author: Cover author ("Prepared by").
        date: Cover date, written as you want it shown.
        organization: Organization name on the cover and in the footer.
        document_type: Small label above the title, e.g. "Quarterly Report".
        logo: Path to a logo image for the cover.
        cover_style: band, full, classic or minimal (default depends on the theme).
        toc: Add a table of contents after the cover.
        page_size: A4, Letter, Legal or A5.
        orientation: portrait or landscape.
        lang: Language code of the text (e.g. "en", "id"): hyphenation, labels, numbers.
        primary_color: Brand color (hex) replacing the theme's primary color.
        accent_color: Brand accent color (hex).
        base_dir: Folder that relative image paths are resolved against (default: output folder).
    """
    result = render_pdf(content, output_path, get_theme(theme, primary_color, accent_color),
                        title, subtitle, author, date, organization=organization,
                        document_type=document_type, logo=logo, cover_style=cover_style,
                        toc=toc, page_size=page_size, orientation=orientation, lang=lang,
                        base_dir=base_dir)
    return (f"PDF created: {result.path} ({result.pages} pages). "
            "Use preview_document to check the layout.")


@mcp.tool(annotations=ToolAnnotations(title="Markdown to Word", read_only_hint=False,
                                     destructive_hint=True, idempotent_hint=False))
@_explain
def markdown_to_docx(content: str, output_path: str | None = None, theme: str = "modern",
                     title: str | None = None, subtitle: str | None = None,
                     author: str | None = None, date: str | None = None,
                     organization: str | None = None, document_type: str | None = None,
                     logo: str | None = None, cover_style: str | None = None,
                     toc: bool = False, page_size: str = "A4", orientation: str = "portrait",
                     lang: str = "en", primary_color: str | None = None,
                     accent_color: str | None = None, base_dir: str | None = None,
                     mode: Literal["create", "append"] = "create",
                     target_path: str | None = None, start_after: str | None = None) -> str:
    """Convert Markdown into a styled, editable Word document.

    Uses real Word styles (headings, lists, captions, TOC) so the file stays easy to edit.
    Same Markdown features and arguments as markdown_to_pdf.

    Args:
        content: Document body in Markdown.
        output_path: Output file path, e.g. ~/Documents/report.docx. In append mode, defaults
            to a sibling .appended.docx file; pass target_path to update target itself with backup.
        theme: modern, corporate, academic, fresh, elegant, minimal, startup or warm.
        title: Adds a cover page with this title.
        subtitle: Cover subtitle.
        author: Cover author ("Prepared by").
        date: Cover date, written as you want it shown.
        organization: Organization name on the cover and in the footer.
        document_type: Small label above the title, e.g. "Proposal".
        logo: Path to a logo image (PNG/JPG) for the cover.
        cover_style: band, full, classic or minimal (default depends on the theme).
        toc: Add a table of contents (Word fills in page numbers when the file opens).
        page_size: A4, Letter, Legal or A5.
        orientation: portrait or landscape.
        lang: Language code of the text (e.g. "en", "id"): proofing language and labels.
        primary_color: Brand color (hex) replacing the theme's primary color.
        accent_color: Brand accent color (hex).
        base_dir: Folder that relative image paths are resolved against (default: output folder).
        mode: "create" for a designed document, or "append" to add content to target_path.
        target_path: Existing .docx to preserve and append into when mode="append".
        start_after: Unique paragraph text or p:index insertion point in append mode.
    """
    before_counts = _docx_counts(target_path) if mode == "append" and target_path else None
    path = render_docx(content, output_path, get_theme(theme, primary_color, accent_color),
                       title, subtitle, author, date, organization=organization,
                       document_type=document_type, logo=logo, cover_style=cover_style,
                       toc=toc, page_size=page_size, orientation=orientation, lang=lang,
                       base_dir=base_dir, mode=mode, target_path=target_path,
                       start_after=start_after)
    note = " Word will offer to update fields on open; accept to fill in the TOC." if toc else ""
    action = "DOCX updated" if mode == "append" else "DOCX created"
    diff = ""
    if before_counts:
        after_paragraphs, after_tables = _docx_counts(str(path))
        added = max(after_paragraphs - before_counts[0], 0)
        removed = max(before_counts[0] - after_paragraphs, 0)
        table_delta = after_tables - before_counts[1]
        diff = f" Diff: {added} paragraph(s) added, {removed} removed; " \
               f"table count changed by {table_delta:+d}."
    return f"{action}: {path}.{diff}{note}"


@mcp.tool(annotations=ToolAnnotations(title="Create Excel workbook", **_WRITE))
@_explain
def create_excel(output_path: str, sheets: list[Sheet], theme: str = "modern",
                 title: str | None = None, author: str | None = None,
                 company: str | None = None, currency: str | None = None, lang: str = "en",
                 primary_color: str | None = None, accent_color: str | None = None) -> str:
    """Create a designed Excel workbook from a declarative description.

    Each sheet can have a banner title, KPI cards, an assumptions block (named input cells),
    Excel tables with formulas, totals, dropdowns, conditional formatting and sparklines,
    native charts and notes. Fonts and colors follow the theme. Formulas are checked
    before writing. Call design_guide(topic="excel") for a full example.

    Args:
        output_path: Output file path, e.g. ~/Documents/sales.xlsx.
        sheets: Sheet definitions (see the schema and design_guide).
        theme: modern, corporate, academic, fresh, elegant, minimal, startup or warm.
        title: Workbook title (document properties).
        author: Workbook author.
        company: Company name (document properties).
        currency: Currency symbol for currency columns, e.g. "Rp", "$", "€".
            Default: "Rp" when lang is "id", otherwise "$".
        lang: Language code, e.g. "en" or "id" (labels and default currency).
        primary_color: Brand color (hex) replacing the theme's primary color.
        accent_color: Brand accent color (hex).
    """
    result = render_xlsx(output_path, sheets, get_theme(theme, primary_color, accent_color),
                         title=title, author=author, company=company, currency=currency,
                         lang=lang)
    lines = [f"Excel workbook created: {result.path}", *result.summary]
    if result.warnings:
        lines += ["Warnings:", *(f"  - {w}" for w in result.warnings)]
    lines.append("Formulas recalculate when the file is opened in Excel.")
    return "\n".join(lines)


@mcp.tool(annotations=ToolAnnotations(title="Preview document", read_only_hint=True))
@_explain
def preview_document(path: str, pages: str = "1-2", width: int = 1000,
                     dpi: int | None = None, format: Literal["png", "jpeg"] = "png",
                     max_pages: int = 6) -> list:
    """Render selected PDF/Office pages as image content plus JSON artifact metadata.

    Use it after creating a document to check covers, page breaks, tables and charts.

    Args:
        path: File to preview.
        pages: Pages to render, e.g. "1", "1,3-5" or "all".
        width: Image width in pixels (300-1800); ignored when dpi is supplied.
        dpi: Render resolution (72-600 DPI), optional.
        format: PNG or JPEG image output.
        max_pages: Limit rendered pages to 1-50 (default 6).
    """
    result = render_pages(path, pages, width, dpi, format, max_pages)
    if not result.images:
        raise ToolError(f"Preview produced no images for {path}; no success response returned.")
    metadata = {
        "path": path,
        "total_pages": result.total_pages,
        "rendered_pages": [image.page for image in result.images],
        "renderer_used": result.renderer_used,
        "note": result.note,
        "images": [{"page": image.page, "path": image.path, "uri": image.uri,
                    "width": image.width, "height": image.height, "format": image.format,
                    "byte_size": len(image.data)} for image in result.images],
    }
    return [json.dumps(metadata, ensure_ascii=False),
            *(Image(data=image.data, format=image.format) for image in result.images)]


def main() -> None:
    transport = os.environ.get("MCP_TRANSPORT", "stdio").strip().lower()
    if transport == "stdio":
        mcp.run()
        return
    if transport != "streamable-http":
        raise ValueError("MCP_TRANSPORT must be 'stdio' or 'streamable-http'")

    allowed_hosts = [value.strip() for value in
                     os.environ.get("MCP_ALLOWED_HOSTS",
                                    "localhost:*,127.0.0.1:*").split(",")
                     if value.strip()]
    allowed_origins = [value.strip() for value in
                       os.environ.get("MCP_ALLOWED_ORIGINS", "").split(",")
                       if value.strip()]
    security = TransportSecuritySettings(allowed_hosts=allowed_hosts,
                                         allowed_origins=allowed_origins)
    mcp.run(
        transport="streamable-http",
        host=os.environ.get("MCP_HOST", "0.0.0.0"),
        port=int(os.environ.get("MCP_PORT", "8000")),
        streamable_http_path=os.environ.get("MCP_PATH", "/mcp"),
        json_response=True,
        stateless_http=True,
        transport_security=security,
    )


if __name__ == "__main__":
    main()
