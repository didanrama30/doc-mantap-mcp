"""Read DOCX structure without requiring an office-suite installation."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import _Cell
from docx.text.paragraph import Paragraph


def _paragraph(paragraph, index: int) -> dict:
    runs = []
    for run in paragraph.runs:
        color = run.font.color.rgb
        runs.append({
            "text": run.text,
            "bold": run.bold,
            "italic": run.italic,
            "underline": run.underline,
            "font": run.font.name,
            "size_pt": run.font.size.pt if run.font.size else None,
            "color": str(color) if color else None,
        })
    ppr = paragraph._p.pPr
    outline = ppr.find(qn("w:outlineLvl")) if ppr is not None else None
    numpr = ppr.find(qn("w:numPr")) if ppr is not None else None
    numbering = None
    if numpr is not None:
        num_id, ilvl = numpr.find(qn("w:numId")), numpr.find(qn("w:ilvl"))
        numbering = {"num_id": num_id.get(qn("w:val")) if num_id is not None else None,
                     "level": ilvl.get(qn("w:val")) if ilvl is not None else None}
    return {
        "index": index,
        "text": paragraph.text,
        "style": paragraph.style.name if paragraph.style else None,
        "alignment": str(paragraph.alignment) if paragraph.alignment is not None else None,
        "outline_level": int(outline.get(qn("w:val"))) if outline is not None else None,
        "numbering": numbering,
        "has_image": bool(paragraph._p.xpath('.//w:drawing | .//w:pict')),
        "runs": runs,
    }


def _cell_text(cell) -> str:
    return "\n".join(p.text for p in cell.paragraphs)


def _table(table, index: int) -> dict:
    rows = []
    max_cols = 0
    for r, row in enumerate(table._tbl.tr_lst):
        cells = []
        logical_col = 0
        for tc in row.tc_lst:
            cell = _Cell(tc, table)
            tcpr = tc.tcPr
            span = tcpr.find(qn("w:gridSpan")) if tcpr is not None else None
            merge = tcpr.find(qn("w:vMerge")) if tcpr is not None else None
            col_span = int(span.get(qn("w:val"))) if span is not None else 1
            cells.append({
                "row": r, "col": logical_col, "text": _cell_text(cell),
                "col_span": col_span,
                "vertical_merge": (merge.get(qn("w:val")) or "continue")
                if merge is not None else None,
            })
            logical_col += col_span
        max_cols = max(max_cols, logical_col)
        rows.append(cells)
    return {"index": index, "rows": len(rows), "cols": max_cols,
            "cells": rows}


def _container_paragraphs(container) -> list[dict]:
    return [_paragraph(p, i) for i, p in enumerate(container.paragraphs)]


def _all_text(doc) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        parts.extend(_cell_text(cell) for row in table.rows for cell in row.cells)
    for section in doc.sections:
        for area in (section.header, section.footer, section.first_page_header,
                     section.first_page_footer, section.even_page_header,
                     section.even_page_footer):
            parts.extend(p.text for p in area.paragraphs)
            for table in area.tables:
                parts.extend(_cell_text(cell) for row in table.rows for cell in row.cells)
    return "\n".join(parts)


def read_docx(path: str, include_runs: bool = True,
              include_headers_footers: bool = True) -> dict:
    """Return document structure and formatting as JSON-serializable data."""
    src = Path(path).expanduser().resolve()
    if not src.is_file():
        raise ValueError(f"DOCX file not found: {src}")
    doc = Document(str(src))
    body_paragraphs = [Paragraph(element, doc) for element in doc.element.body.iter(qn("w:p"))]
    paragraphs = [_paragraph(p, i) for i, p in enumerate(body_paragraphs)]
    tables = [_table(table, i) for i, table in enumerate(doc.tables)]
    blocks = []
    for child in doc.element.body:
        if child.tag == qn("w:p"):
            text = "".join(child.xpath('.//w:t/text()'))
            blocks.append({"type": "paragraph", "index": len(blocks), "text": text})
        elif child.tag == qn("w:tbl"):
            blocks.append({"type": "table", "index": len(blocks),
                           "table_index": sum(b["type"] == "table" for b in blocks)})
    headers_footers = []
    if include_headers_footers:
        for i, section in enumerate(doc.sections):
            for label, area in (
                ("header", section.header), ("footer", section.footer),
                ("first_page_header", section.first_page_header),
                ("first_page_footer", section.first_page_footer),
                ("even_page_header", section.even_page_header),
                ("even_page_footer", section.even_page_footer),
            ):
                headers_footers.append({"section": i, "kind": label,
                                        "paragraphs": _container_paragraphs(area),
                                        "tables": [_table(t, n) for n, t in enumerate(area.tables)]})
    if not include_runs:
        for item in paragraphs:
            item.pop("runs", None)
    all_text = _all_text(doc)
    placeholders = sorted(set(re.findall(r"\{\{\s*[^{}]+\s*}}|<<[^<>]+>>", all_text)))
    sections = [{
        "page_width_inches": round(s.page_width.inches, 3),
        "page_height_inches": round(s.page_height.inches, 3),
        "orientation": "landscape" if s.page_width > s.page_height else "portrait",
        "margins_inches": {name: round(getattr(s, f"{name}_margin").inches, 3)
                            for name in ("top", "right", "bottom", "left")},
    } for s in doc.sections]
    return {
        "path": str(src),
        "metadata": {"title": doc.core_properties.title, "author": doc.core_properties.author,
                     "subject": doc.core_properties.subject, "sections": sections},
        "paragraphs": paragraphs,
        "tables": tables,
        "blocks": blocks,
        "headers_footers": headers_footers,
        "placeholders": placeholders,
    }


def outline_docx(path: str) -> dict:
    """Return a compact heading/table outline for agent context."""
    data = read_docx(path, include_runs=False, include_headers_footers=False)
    headings = [{"index": p["index"], "text": p["text"], "style": p["style"]}
                for p in data["paragraphs"] if (p["style"] or "").lower().startswith("heading")]
    tables = [{"index": t["index"], "rows": t["rows"], "cols": t["cols"],
               "first_row": [c["text"] for c in (t["cells"][0] if t["cells"] else [])]}
              for t in data["tables"]]
    return {"path": data["path"], "headings": headings, "tables": tables,
            "placeholders": data["placeholders"]}
